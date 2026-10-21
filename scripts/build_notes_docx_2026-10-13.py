"""Build a clean, structured Word document from the 2026-10-13 meeting notes.

Same visual base as build_notes_docx.py (2026-09-09 notes): identical styles,
colors and helper functions, reused as-is for consistency.
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x26, 0x46, 0x53)      # deep teal
ACCENT_HEX = "264653"
LIGHT_HEX = "E8EEF0"                      # pale teal fill
CALLOUT_HEX = "F3EDE2"                    # warm sand fill
GREY = RGBColor(0x6B, 0x6B, 0x6B)
DARK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Pt(64)
sec.top_margin = sec.bottom_margin = Pt(56)

# --- base styles -----------------------------------------------------------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.12

for lvl, size in ((1, 15), (2, 11.5)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = ACCENT
    h.paragraph_format.space_before = Pt(18 if lvl == 1 else 10)
    h.paragraph_format.space_after = Pt(5)
    h.paragraph_format.keep_with_next = True

for s in ("List Bullet", "List Bullet 2", "List Number"):
    st = doc.styles[s]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = DARK
    st.paragraph_format.space_after = Pt(3)


def _shade(el, hex_fill):
    tcpr = el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def bottom_border(paragraph, hex_color=ACCENT_HEX, sz="8"):
    p = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), hex_color)
    borders.append(bottom)
    p.append(borders)


def para(text, size=None, color=None, italic=False, bold=False,
         align=None, space_before=None, space_after=None):
    par = doc.add_paragraph()
    run = par.add_run(text)
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    run.font.italic = italic
    run.font.bold = bold
    if align is not None:
        par.alignment = align
    if space_before is not None:
        par.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        par.paragraph_format.space_after = Pt(space_after)
    return par


def bullets(items, style="List Bullet"):
    for it in items:
        if isinstance(it, tuple):
            para_bullet(it[0], style)
            bullets(it[1], style="List Bullet 2")
        else:
            para_bullet(it, style)


def para_bullet(text, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def callout(label, text, fill=CALLOUT_HEX):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    _shade(cell._tc, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + " ")
    r.font.bold = True
    r.font.color.rgb = ACCENT
    r.font.size = Pt(10.5)
    r2 = p.add_run(text)
    r2.font.size = Pt(10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def data_table(headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    for i, h in enumerate(headers):
        _shade(hdr[i]._tc, ACCENT_HEX)
        p = hdr[i].paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        cells = tbl.add_row().cells
        for ci, val in enumerate(row):
            if ri % 2 == 1:
                _shade(cells[ci]._tc, LIGHT_HEX)
            p = cells[ci].paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(10)
            if ci == 0:
                r.font.bold = True
                r.font.color.rgb = ACCENT
    if widths:
        for row in tbl.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Pt(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def checklist(items):
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    for it in items:
        cells = tbl.add_row().cells
        cells[0].width = Pt(20)
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rb = p0.add_run("\u2610")
        rb.font.size = Pt(12)
        rb.font.color.rgb = ACCENT
        r = cells[1].paragraphs[0].add_run(it)
        r.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


# --- title ---------------------------------------------------------------
para("Notes de réunion", size=24, bold=True, color=ACCENT, space_after=2)
t = para("PYB · IB  —  13 octobre 2026", size=11, color=GREY, space_after=12)
bottom_border(t)
para("Projet PIPELINE_01 — points abordés et décisions.",
     italic=True, color=GREY, space_after=14)

# --- à faire ---------------------------------------------------------------
doc.add_heading("À faire", level=1)
checklist([
    "QC des images dixon : passer l'affichage en 5 lignes et 4 colonnes.",
    "Pré-remplir l'interface graphique : dès qu'une section est remplie par l'utilisateur, "
    "pré-remplir les autres. Avoir in fine les 4 sections du rapport disponibles facilement.",
    "Privilégier le mode de fonctionnement optimal ; faire abstraction de la GUI, qui est "
    "un moyen pour que les techs utilisent au plus tôt l'outil, pas une fin en soi — les "
    "contraintes de la GUI ne doivent surtout pas influencer le fonctionnement du programme.",
    "Pouvoir aller chercher les informations via le PACS.",
    "Vérifier l'erreur « where without out » du T2.",
])

# --- prochain chantier : validation des méthodes ---------------------------
doc.add_heading("Prochain chantier — Validation des méthodes, qualité des résultats",
                 level=1)
para("Décrire les méthodes et documenter les tests.")
bullets([
    "Comparer nos résultats avec ceux de Matlab.",
    "Vérification manuelle : ouvrir les cartographies dixon et les ROIs dans ITK-SNAP, "
    "comparer avec celles générées par notre programme.",
    "Définir une marge de tolérance à l'avance et vérifier si les résultats s'y trouvent.",
    "Mettre en place des méthodes de vérification des procédés mis en œuvre, et les "
    "documenter.",
    "Rapport de validation : à telle date, tels tests ont été faits — signé et daté.",
])

doc.add_heading("Où est-ce que les erreurs peuvent arriver ?", level=2)
callout("Constat.",
        "Le plus souvent, ce sont des erreurs humaines : mauvaise date, mauvaises séries, "
        "etc.")

# --- debug -------------------------------------------------------------------
doc.add_heading("Debug", level=1)
bullets([
    "Pouvoir identifier les sources.",
    "Savoir quand ça a été fait, par qui, avec quelle version du code — une signature "
    "unique pour reproduire exactement la même chose si besoin.",
    "À réfléchir : si on veut faire plus compliqué pour mutualiser des éléments.",
])

# --- segmentation ------------------------------------------------------------
doc.add_heading("Côté segmentation", level=1)
callout("Cas à traiter.",
        "Avec l'outil de Louis « Click and Fix », utilisé pour une section : éviter de "
        "refaire le travail pour une autre section similaire. Réfléchir à ne pas perdre ce "
        "travail.")
bullets([
    "Penser au stockage de ces fichiers précieux.",
    "Réfléchir à si on veut faire quelque chose de plus compliqué pour mutualiser les "
    "éléments.",
])

# --- à proposer ---------------------------------------------------------------
doc.add_heading("À proposer", level=1)
bullets([
    "Lien entre méthode et segmentation.",
    "Découpler la segmentation de dixon ?",
    "Segmentation unique pour tous les process du même segment.",
    "Pouvoir fournir sa propre segmentation.",
    "Comment coupler proprement la segmentation d'une méthode à l'entrée d'une autre ?",
])

out = r"c:\Users\i.benziane\Documents\PIPELINE_01\Notes_reunion_2026-10-13.docx"
doc.save(out)
print("saved:", out)
