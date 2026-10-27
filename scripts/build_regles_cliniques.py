"""Build MuReport-regles-cliniques.docx: clinical rules coded in muReport, for physician validation.

Same visual base as build_rapport_tests_auto001.py.
"""


from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x26, 0x46, 0x53)      # deep teal
ACCENT_HEX = "264653"
LIGHT_HEX = "E8EEF0"                      # pale teal fill
CALLOUT_HEX = "F3EDE2"                    # warm sand fill
GREY = RGBColor(0x6B, 0x6B, 0x6B)
DARK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Pt(64)
sec.top_margin = sec.bottom_margin = Pt(56)

# --- base styles -----------------------------------------------------------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.12

for lvl, size in ((1, 15), (2, 11.5)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = ACCENT
    h.paragraph_format.space_before = Pt(18 if lvl == 1 else 10)
    h.paragraph_format.space_after = Pt(5)
    h.paragraph_format.keep_with_next = True

for s in ("List Bullet", "List Bullet 2", "List Number"):
    st = doc.styles[s]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = DARK
    st.paragraph_format.space_after = Pt(3)


def _shade(el, hex_fill):
    tcpr = el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def bottom_border(paragraph, hex_color=ACCENT_HEX, sz="8"):
    p = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), hex_color)
    borders.append(bottom)
    p.append(borders)


def para(text, size=None, color=None, italic=False, bold=False,
         align=None, space_before=None, space_after=None):
    par = doc.add_paragraph()
    run = par.add_run(text)
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    run.font.italic = italic
    run.font.bold = bold
    if align is not None:
        par.alignment = align
    if space_before is not None:
        par.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        par.paragraph_format.space_after = Pt(space_after)
    return par


def bullets(items, style="List Bullet"):
    for it in items:
        if isinstance(it, tuple):
            para_bullet(it[0], style)
            bullets(it[1], style="List Bullet 2")
        else:
            para_bullet(it, style)


def para_bullet(text, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def callout(label, text, fill=CALLOUT_HEX):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    _shade(cell._tc, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + " ")
    r.font.bold = True
    r.font.color.rgb = ACCENT
    r.font.size = Pt(10.5)
    r2 = p.add_run(text)
    r2.font.size = Pt(10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def data_table(headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    for i, h in enumerate(headers):
        _shade(hdr[i]._tc, ACCENT_HEX)
        p = hdr[i].paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        cells = tbl.add_row().cells
        for ci, val in enumerate(row):
            if ri % 2 == 1:
                _shade(cells[ci]._tc, LIGHT_HEX)
            p = cells[ci].paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(10)
            if ci == 0:
                r.font.bold = True
                r.font.color.rgb = ACCENT
    if widths:
        for row in tbl.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Pt(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def checklist(items):
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    for it in items:
        cells = tbl.add_row().cells
        cells[0].width = Pt(20)
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rb = p0.add_run("\u2610")
        rb.font.size = Pt(12)
        rb.font.color.rgb = ACCENT
        r = cells[1].paragraphs[0].add_run(it)
        r.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def rule(number, title, what, where, question):
    doc.add_heading(f"{number}. {title}", level=1)
    para(what)
    p = para(f"Dans le code : {where}", size=9, color=GREY, italic=True)
    callout("À valider :", question)


# --- title ---------------------------------------------------------------
para("Règles cliniques codées dans muReport", size=22, bold=True, color=ACCENT, space_after=2)
t = para("muReport · Inès Benziane  —  22 octobre 2026  —  à faire valider par un médecin",
         size=11, color=GREY, space_after=12)
bottom_border(t)
para("Le rapport médical contient des phrases et des indicateurs choisis automatiquement à partir "
     "de seuils écrits dans le code. Ce document les rend visibles pour qu'un médecin les valide, "
     "les corrige ou les refuse. Chaque règle est vérifiée telle quelle par les tests automatisés "
     "(AUTO-001) : si une règle change, le test correspondant doit changer avec elle.",
     italic=True, color=GREY, space_after=14)

rule(1, "Seuil de T2 « élevé » : 39 ms",
     "Un muscle est considéré comme ayant un T2 élevé si sa moyenne de T2 est supérieure ou égale à "
     "39 ms (39,0 ms compte comme élevé, 38,99 ms non). Ce seuil sert au commentaire T2 (règle 2) et "
     "au signal de franchissement de seuil (règle 5).",
     "comment_generator/config/t2_comments_fr.yaml et _en.yaml (elevated_threshold_ms), "
     "synthesis_builder.py (T2_THRESHOLD). Un test vérifie que ces trois valeurs restent identiques.",
     "La valeur de 39 ms, et le fait que l'égalité compte comme « élevé ».")

doc.add_heading("2. Commentaire T2 selon le nombre de muscles élevés", level=1)
para("Le rapport affiche une phrase choisie selon le nombre de muscles dont le T2 est ≥ 39 ms :")
data_table(["Muscles ≥ 39 ms", "Phrase affichée (version française)"], [
    ["0", "« Valeurs des moyennes du T2 musculaire dans la norme. »"],
    ["1", "« Discrète augmentation de la moyenne du T2 musculaire en regard de [nom du muscle] "
          "suggérant l'existence de processus pathologiques actifs. »"],
    ["2 ou 3", "« Augmentation modérée et significative des moyennes du T2 musculaire dans plusieurs "
               "territoires musculaires suggérant l'existence de processus pathologiques actifs. »"],
    ["4 ou plus", "« Augmentation significative des moyennes du T2 musculaire dans plusieurs territoires "
                  "musculaires suggérant l'existence de processus pathologiques actifs. »"],
], widths=[90, 370])
para("Dans le code : comment_generator/selector.py, phrases dans t2_comments_fr.yaml et t2_comments_en.yaml.",
     size=9, color=GREY, italic=True)
callout("À valider :", "Les quatre paliers (0, 1, 2-3, 4 et plus), et le texte de chaque phrase en français et en anglais.")

rule(3, "Phrase « normale limite » jamais affichée",
     "Une cinquième phrase est prévue : « Valeurs de T2 dans la norme, suggérant que les processus "
     "pathologiques sont peu actifs. » Sa condition d'affichage est « aucun muscle ≥ 39 ms et T2 "
     "maximal ≥ 39 ms », ce qui est impossible. Elle n'apparaît donc dans aucun rapport. "
     "L'intention était peut-être : aucun muscle au-dessus du seuil, mais au moins un proche du seuil.",
     "t2_comments_*.yaml (catégorie normal_borderline). Incident ouvert, test en échec attendu.",
     "Cette phrase est-elle souhaitée ? Si oui, à partir de quelle valeur de T2 (par exemple entre 37 et "
     "39 ms) ? Sinon, la supprimer.")

doc.add_heading("4. Évolution entre deux examens (page de synthèse)", level=1)
para("Pour chaque muscle, la valeur actuelle est comparée à celle de l'examen antérieur :")
bullets([
    "T2 : différence en millisecondes ; FF : différence en points de pourcentage (par exemple 20 % → 30 % = +10 points).",
    "La différence est arrondie à l'entier, puis comparée au seuil : « aggravé » si elle dépasse "
    "strictement +1 ms (T2) ou +5 points (FF), « amélioré » si elle est en dessous de −1 ms ou −5 points, "
    "« stable » sinon.",
    "Si l'un des deux examens n'a pas de valeur pour ce muscle : « inconnu ».",
])
data_table(["Exemple (T2)", "Différence", "Après arrondi", "Statut affiché"], [
    ["40,0 → 41,4 ms", "+1,4 ms", "+1", "stable"],
    ["40,0 → 41,6 ms", "+1,6 ms", "+2", "aggravé"],
    ["40,0 → 38,0 ms", "−2,0 ms", "−2", "amélioré"],
], widths=[110, 100, 100, 150])
data_table(["Exemple (FF)", "Différence", "Après arrondi", "Statut affiché"], [
    ["20 % → 25 %", "+5 points", "+5", "stable"],
    ["20 % → 30 %", "+10 points", "+10", "aggravé"],
], widths=[110, 100, 100, 150])
para("Dans le code : section_generator/src/domain/services/synthesis_builder.py (EVOLUTION_THRESHOLD, compute_evolution).",
     size=9, color=GREY, italic=True)
callout("À valider :", "Les seuils (1 ms pour le T2, 5 points pour la FF), l'arrondi à l'entier avant comparaison, "
        "et le fait qu'une variation égale au seuil soit classée « stable ».")

rule(5, "Signal de franchissement du seuil T2",
     "Un muscle est signalé quand son T2 passe sous le seuil à l'examen antérieur et au-dessus ou au "
     "niveau du seuil à l'examen actuel (antérieur < 39 ms ≤ actuel). Uniquement pour le T2.",
     "synthesis_builder.py (threshold_crossed).",
     "L'utilité de ce signal, et s'il faut aussi signaler le passage inverse (retour sous le seuil).")

rule(6, "Muscles mis en avant dans la synthèse",
     "Les muscles aggravés et améliorés sont classés du plus grand au plus petit changement. Les trois "
     "muscles ayant le plus grand changement, dans un sens ou dans l'autre, sont mis en avant.",
     "synthesis_builder.py (top_muscles).",
     "Le nombre de muscles mis en avant (3), et le fait de mélanger aggravations et améliorations.")

rule(7, "Muscles affichés dans les tableaux de valeurs (cuisses)",
     "Pour les cuisses, la segmentation détoure 13 muscles, mais les tableaux de valeurs du rapport n'en "
     "affichent que 11 : le court adducteur (AB) et le chef court du biceps fémoral (BF_SH) n'y figurent pas. "
     "Leurs valeurs sont calculées mais visibles seulement au survol du schéma, ce qui est impossible sur un "
     "PDF imprimé. Les ajouter modifie la mise en page du rapport.",
     "gabarits templates/FF/1slice/thighs et templates/T2/1slice/thighs (muscles_list), légende anatomique "
     "de report_template.html. Incident ouvert, test en échec attendu.",
     "Les valeurs de AB et BF_SH sont-elles utiles au médecin ? Si oui, les ajouter aux tableaux (évolution de mise en page).")

rule(8, "Voxels retenus pour le T2 moyen : 15 à 70 ms",
     "Pour calculer le T2 moyen d'un muscle, seuls les voxels dont le T2 est compris entre 15 et 70 ms (et dont "
     "l'intervalle de confiance est inférieur ou égal à 10) sont retenus. Pour un muscle très inflammatoire dont "
     "une partie des voxels dépasse 70 ms, ces voxels sont exclus et la moyenne affichée est sous-estimée. "
     "Décision du 22/10/2026 : la borne de 70 ms est conservée pour l'instant.",
     "methods/get_results/config/results_config.yml (t2map_3exp, filters). Vérifié par les tests automatisés.",
     "Proposition à discuter : relever la borne haute, ou signaler dans le rapport la part de voxels exclus "
     "au-dessus de 70 ms.")

# --- validation -------------------------------------------------------------
doc.add_heading("Validation", level=1)
data_table(["Règle", "Décision (validée / à modifier / refusée)", "Commentaire"],
           [[str(i), "", ""] for i in range(1, 9)], widths=[50, 170, 240])
para("Nom du médecin : ____________________     Date : ____________     Signature : ____________________",
     space_before=10)

out = "MuReport-regles-cliniques.docx"
doc.save(out)
print("saved", out)


