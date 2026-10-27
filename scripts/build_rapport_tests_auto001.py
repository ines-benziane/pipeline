"""Build the AUTO-001 test report (readable proof for an auditor).

Same visual base as build_notes_docx_2026-10-13.py: styles and helpers reused as-is.
Output: MuReport-rapport-tests-AUTO-001.docx (copied to doc/MuReport).
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x26, 0x46, 0x53)      # deep teal
ACCENT_HEX = "264653"
LIGHT_HEX = "E8EEF0"                      # pale teal fill
CALLOUT_HEX = "F3EDE2"                    # warm sand fill
GREY = RGBColor(0x6B, 0x6B, 0x6B)
DARK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Pt(64)
sec.top_margin = sec.bottom_margin = Pt(56)

# --- base styles -----------------------------------------------------------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.12

for lvl, size in ((1, 15), (2, 11.5)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = ACCENT
    h.paragraph_format.space_before = Pt(18 if lvl == 1 else 10)
    h.paragraph_format.space_after = Pt(5)
    h.paragraph_format.keep_with_next = True

for s in ("List Bullet", "List Bullet 2", "List Number"):
    st = doc.styles[s]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = DARK
    st.paragraph_format.space_after = Pt(3)


def _shade(el, hex_fill):
    tcpr = el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def bottom_border(paragraph, hex_color=ACCENT_HEX, sz="8"):
    p = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), hex_color)
    borders.append(bottom)
    p.append(borders)


def para(text, size=None, color=None, italic=False, bold=False,
         align=None, space_before=None, space_after=None):
    par = doc.add_paragraph()
    run = par.add_run(text)
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    run.font.italic = italic
    run.font.bold = bold
    if align is not None:
        par.alignment = align
    if space_before is not None:
        par.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        par.paragraph_format.space_after = Pt(space_after)
    return par


def bullets(items, style="List Bullet"):
    for it in items:
        if isinstance(it, tuple):
            para_bullet(it[0], style)
            bullets(it[1], style="List Bullet 2")
        else:
            para_bullet(it, style)


def para_bullet(text, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def callout(label, text, fill=CALLOUT_HEX):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    _shade(cell._tc, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + " ")
    r.font.bold = True
    r.font.color.rgb = ACCENT
    r.font.size = Pt(10.5)
    r2 = p.add_run(text)
    r2.font.size = Pt(10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def data_table(headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    for i, h in enumerate(headers):
        _shade(hdr[i]._tc, ACCENT_HEX)
        p = hdr[i].paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        cells = tbl.add_row().cells
        for ci, val in enumerate(row):
            if ri % 2 == 1:
                _shade(cells[ci]._tc, LIGHT_HEX)
            p = cells[ci].paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(10)
            if ci == 0:
                r.font.bold = True
                r.font.color.rgb = ACCENT
    if widths:
        for row in tbl.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Pt(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def checklist(items):
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    for it in items:
        cells = tbl.add_row().cells
        cells[0].width = Pt(20)
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rb = p0.add_run("\u2610")
        rb.font.size = Pt(12)
        rb.font.color.rgb = ACCENT
        r = cells[1].paragraphs[0].add_run(it)
        r.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


# --- title ---------------------------------------------------------------
para("Rapport de tests automatisés — AUTO-001", size=22, bold=True, color=ACCENT, space_after=2)
t = para("muReport · Inès Benziane  —  22 octobre 2026  —  version du code 6e33e57",
         size=11, color=GREY, space_after=12)
bottom_border(t)
para("Ce document décrit ce qui a été testé automatiquement, comment, et avec quel résultat. "
     "Il accompagne MuReport-plan-de-tests.xlsx (ligne AUTO-001, onglets Journal et Rapports "
     "d'incidents) et applique la stratégie de MuReport-politique-tests.docx.",
     italic=True, color=GREY, space_after=14)

# --- 1. synthèse -----------------------------------------------------------
doc.add_heading("1. Synthèse", level=1)
callout("Verdict : conforme.",
        "231 cas de test sur les 3 briques livrées (muReport, medical_report, dicom), dont des tests de bout en bout "
        "jusqu'au rapport : 225 réussis, 6 échecs attendus (anomalies connues, voir section 6). "
        "Exécution du 22/10/2026 sur la station Prométhée, commit 6e33e57.")
data_table(["Élément", "Valeur"], [
    ["Identifiant", "AUTO-001 (plan de tests, phase 1)"],
    ["Niveau", "Tests unitaires et tests d'intégration"],
    ["Outil", "pytest 9.1.1, Python 3.12.13"],
    ["Environnement", "Station Prométhée (Linux), environnement conda « pipeline », Docker (segmentation MuSeg-AI)"],
    ["Version du code", "6e33e57 (dépôt github.com/ines-benziane/pipeline), avec medical_report 1725584 et dicom 0e43e4f"],
    ["Durée", "4 min 12 s"],
    ["Preuve d'exécution", "resultats_tests/AUTO-001_20261022_6e33e57.txt (+ .xml), voir section 8"],
], widths=[130, 330])

# --- 2. périmètre ----------------------------------------------------------
doc.add_heading("2. Périmètre", level=1)
para("Inclus : les trois briques livrées, des DICOM jusqu'au rapport, hors calcul.")
bullets([
    "saisies utilisateur (numéros de séries, acquisition, méthode) ;",
    "orchestration : création, exécution, pause pour contrôle qualité (QC), reprise et sauvegarde des jobs ;",
    "méthodes dixon3pt et t2map_3exp : appel par le moteur et contrôles d'entrée ;",
    "chaîne complète, des DICOM aux valeurs par muscle, pour la méthode Dixon ;",
    "écriture des résultats (tableau mutools vers JSON) et génération du rapport ;",
    "ligne de commande, images de contrôle qualité, extraction des valeurs par muscle ;",
    "medical_report : lecture des résultats, choix des versions et des examens antérieurs, "
    "commentaire T2, synthèse d'évolution ;",
    "dicom : anonymisation et pseudonymisation, requête au PACS, sélection des séries, "
    "tri des fichiers reçus, liste de patients.",
])
para("Exclus, et pourquoi :")
bullets([
    "le calcul lui-même (mutools) : logique de recherche publiée, non modifiable par muReport. "
    "Sa justesse est vérifiée contre Matlab par les tests RES-001 et RES-002 ;",
    "les échanges réels avec le PACS Telemis (C-FIND, C-MOVE) : non automatisables sans serveur, "
    "vérifiés lors des transferts réels ; la requête envoyée et le traitement des fichiers reçus sont testés ;",
    "l'interface graphique : couverte par les tests FONC-001 et VAL (utilisation réelle par les techniciens).",
])

# --- 3. méthode ------------------------------------------------------------
doc.add_heading("3. Méthode", level=1)
bullets([
    "Tests unitaires : chaque fonction est vérifiée isolément. Les méthodes de calcul sont remplacées "
    "par des méthodes factices, sans image ni calcul : les tests sont rapides et donnent toujours le même résultat.",
    "Tests d'intégration : plusieurs modules enchaînés. Pipeline complet avec une méthode factice, "
    "et méthode Dixon réelle sur un examen réel.",
    "Test d'équivalence avec mutools : un examen de référence (CL2LD260112, jambes, séries 17 et 18) "
    "est traité par muReport, et les valeurs de chacun des 16 muscles (nombre de pixels NPIX, fraction "
    "de graisse FF) sont comparées à celles obtenues avec mutools utilisé seul. Ce test prouve que muReport "
    "ne déforme pas les résultats de mutools (bonnes séries lues, bonne orientation, bonne ROI, bonnes "
    "valeurs associées à chaque muscle). Il ne prouve pas la justesse clinique, objet de RES-001 : "
    "l'examen de référence présente d'ailleurs une inversion eau/graisse connue sur la jambe droite, "
    "reproduite à l'identique.",
    "Test de bout en bout : les fichiers de résultats d'un vrai examen (anonymisé) sont transformés en rapport "
    "avec la configuration construite par l'interface graphique ; chaque valeur affichée est comparée à celle "
    "du fichier, pour le bon muscle et le bon côté. Le contrôle porte sur le HTML converti en PDF.",
    "Isolation : les tests écrivent uniquement dans des dossiers temporaires, jamais dans les données réelles.",
])
doc.add_heading("Tolérances", level=2)
data_table(["Grandeur", "Tolérance", "Justification"], [
    ["FF par muscle", "± 0,1 %", "Biomarqueur clinique. Écart maximal observé : 0,016 %."],
    ["NPIX par muscle", "± 0,1 % ou ± 10 pixels (le plus large)",
     "La segmentation MuSeg-AI n'est pas reproductible au pixel près : la bordure des muscles varie de "
     "1 à 5 pixels d'une exécution à l'autre (3 exécutions, maximum 0,056 %). Un vrai défaut "
     "(image retournée, ROI décalée) modifie des milliers de pixels."],
    ["Autres tests", "Aucune", "Résultat exact attendu."],
], widths=[90, 120, 250])

# --- 4. détail -------------------------------------------------------------
doc.add_heading("4. Détail des tests", level=1)
data_table(["Domaine", "Ce qui est vérifié", "Cas", "Résultat"], [
    ["Saisies utilisateur", "Formats valides correctement convertis ; formats invalides (« 17,,18 », "
     "« abc », acquisition à plus de 3 parties, paramètre sans valeur) rejetés avec une erreur, jamais devinés.",
     "18", "18/18"],
    ["Job et sauvegarde", "Identifiant unique, mode QC valide, état initial ; un job sauvegardé puis "
     "rechargé est identique (condition de la reprise après QC ou après crash).", "10", "10/10"],
    ["Exécution d'un job", "Succès → résultat publié ; pause QC → job suspendu, étape enregistrée ; "
     "échec → job en échec, aucun résultat publié ; reprise depuis l'étape enregistrée ; dossier QC ; journal du run.",
     "12", "12/12"],
    ["Reprise après QC", "Reprise à l'étape enregistrée, même job, décision de l'utilisateur transmise, "
     "dossier QC conservé, relance complète si une correction est demandée, job inconnu rejeté. "
     "QC refusé → aucun résultat.", "9", "8/8 + 1 échec attendu"],
    ["Méthodes dixon3pt et t2map_3exp", "Acceptent les appels du moteur ; action inconnue, séries absentes "
     "du dossier, étape inconnue, segment inconnu → erreurs explicites.", "11", "11/11"],
    ["Registre et classe de base", "Méthode inconnue rejetée, méthode incomplète refusée, contrôle des "
     "étapes et des actions autorisées.", "11", "11/11"],
    ["Pipeline complet (intégration)", "Résolution de l'examen (aucun, un seul, plusieurs candidats), "
     "transmission des paramètres, traçabilité des séries choisies, pause QC, échec sans résultat.",
     "10", "10/10"],
    ["Catalogue d'examens", "Recherche d'examens, examens liés, liste des séries.", "3", "3/3"],
    ["Équivalence Dixon / mutools (intégration, données réelles)", "16 muscles : NPIX et FF conformes "
     "à la référence mutools dans les tolérances.", "1", "1/1"],
    ["Écriture des résultats (muReport)", "Tableau mutools vers JSON : un muscle par nom et côté, valeurs "
     "inchangées, coupes triées, métadonnées et décision QC transmises, nom du fichier.", "14", "14/14"],
    ["Génération du rapport (muReport)", "Erreur explicite si aucun PDF n'est produit ; jamais d'ancien PDF renvoyé.", "2", "2/2"],
    ["Lecture des résultats (medical_report)", "Examen le plus récent ou à la date demandée ; valeurs et identité "
     "lues sans modification ; jamais un autre patient, segment ou version ; fichier corrompu ignoré.", "8", "8/8"],
    ["Versions et antécédents (medical_report)", "Règles de version (≥, ≤, exacte), méthode dont le nom contient « _ », "
     "version la plus récente (1.10 > 1.9) ; antécédents antérieurs, du même patient, du plus ancien au plus récent.", "12", "12/12"],
    ["Commentaire T2 (medical_report)", "Phrase choisie selon le nombre de muscles ≥ 39 ms, seuil inclusif, "
     "muscle nommé, français et anglais, même seuil partout. Phrase « normale limite » jamais affichée.",
     "12", "11/11 + 1 échec attendu"],
    ["Synthèse d'évolution (medical_report)", "Évolution T2 (ms) et FF (points), arrondi et seuils, valeur manquante, "
     "appariement des muscles, franchissement du seuil T2, classement des muscles.", "20", "20/20"],
    ["Anonymisation et pseudonymisation (dicom)", "Données démographiques supprimées ; nom remplacé par un pseudonyme "
     "stable d'un run à l'autre ; date de naissance et sexe vidés ; table de correspondance sans le nom.", "11", "11/11"],
    ["Requête et sélection des séries (dicom)", "Critères transmis au PACS, informations demandées en retour ; "
     "2 séries les plus anciennes, la plus récente, séries sans date écartées.", "11", "11/11"],
    ["Réception (dicom)", "Noms de dossiers valides, tri par patient puis par série, dossier de transit vidé, "
     "series_metadata.json, lecture de la liste de patients (Excel, CSV).", "10", "10/10"],
    ["Rapport de bout en bout (intégration)", "PDF produit ; date d'examen ; chaque valeur FF et T2 des tableaux "
     "sur le bon muscle et le bon côté (jambes, cuisses) ; commentaire T2 ; évolution par rapport à l'examen "
     "antérieur. Identité du patient et muscles AB, BF_SH : non affichés.", "14", "12/12 + 2 échecs attendus"],
    ["Ligne de commande", "Chaque option transmise au pipeline, valeurs par défaut, option obligatoire manquante "
     "refusée et nommée, messages de fin, erreurs propres ; process et report de bout en bout. Saisie mal formée "
     "et dossier inexistant : message non explicite.", "17", "15/15 + 2 échecs attendus"],
    ["Images QC et extraction par muscle", "Image de contrôle indépendante de l'orientation stockée, une image GIF "
     "par coupe, ROI dessinée ; FF et T2 moyens par muscle, érosion des bords, voxels hors limites exclus "
     "(FF hors [0, 1], T2 hors 15-70 ms).", "15", "15/15"],
    ["Total", "", "231", "225 réussis, 6 échecs attendus"],
], widths=[105, 255, 35, 70])

# --- 5. historique ---------------------------------------------------------
doc.add_heading("5. Historique des exécutions", level=1)
data_table(["Date", "Lieu", "Résultat", "Verdict", "Commit"], [
    ["22/10/2026", "PC (sans Docker)", "83 réussis, 1 échec attendu, 1 non exécuté (équivalence Dixon)", "Pass", "7969096"],
    ["22/10/2026", "Prométhée", "Équivalence Dixon : écart NPIX de 1 à 5 px sur 10 muscles, FF conforme → incident", "Fail", "7969096"],
    ["22/10/2026", "Prométhée", "Après fixation de la tolérance NPIX : 84 réussis, 1 échec attendu", "Pass", "37c3b0b"],
    ["22/10/2026", "Prométhée", "Périmètre étendu aux 3 briques : non exécuté, paquet dicom non installé (environnement)", "Fail", "f67bce4"],
    ["22/10/2026", "Prométhée", "3 briques : 183 réussis, 2 échecs attendus", "Pass", "f67bce4"],
    ["22/10/2026", "Prométhée", "Avec le test de bout en bout : 195 réussis, 4 échecs attendus", "Pass", "b226714"],
    ["22/10/2026", "Prométhée", "Avec CLI, QC et extraction ; code mort déplacé : 225 réussis, 6 échecs attendus", "Pass", "6e33e57"],
], widths=[60, 70, 230, 40, 55])
para("Les exécutions en échec sont conservées volontairement : la première montre que le test détecte un "
     "écart et que la tolérance a été fixée à partir de mesures, pas pour faire passer le test ; la seconde "
     "est un défaut d'installation de l'environnement, corrigé avant la ré-exécution.",
     italic=True, color=GREY)

# --- 6. anomalies ----------------------------------------------------------
doc.add_heading("6. Anomalies détectées par les tests", level=1)
data_table(["Anomalie", "Statut"], [
    ["Reprise après QC impossible avec dixon3pt et t2map_3exp (paramètre seg_series refusé).", "Corrigée (7969096)"],
    ["Mode QC et dossier QC perdus à la sauvegarde d'un job.", "Corrigée (7969096)"],
    ["Saisie d'acquisition à plus de 3 parties acceptée sans erreur.", "Corrigée (7969096)"],
    ["Segmentation MuSeg-AI non reproductible au pixel près (gravité faible, FF non impactée).",
     "Résolue : tolérance NPIX (37c3b0b)"],
    ["QC refusé : un résultat est quand même produit (FONC-004, gravité moyenne). Chemin de reprise en "
     "ligne de commande, non utilisé par l'interface graphique en phase 1.",
     "Ouverte : refonte du QC, phase 2"],
    ["medical_report : la version 1.9 d'une méthode était choisie avant la 1.10 (tri alphabétique).", "Corrigée (9032149)"],
    ["medical_report : la phrase T2 « normale limite » ne s'affiche jamais (condition contradictoire).",
     "Ouverte : à trancher avec un médecin (MuReport-regles-cliniques.docx)"],
    ["dicom : l'option --anonymize est acceptée mais ignorée par les scripts de transfert C-MOVE (gravité élevée). "
     "Trouvé à la lecture du code.", "Ouverte : utiliser la pseudonymisation (-rps), vérifiée par les tests"],
    ["dicom : la pseudonymisation conserve le PatientID (clé de correspondance).", "Ouverte : à confirmer selon la nature du PatientID"],
    ["Rapport : ni nom du patient ni identifiant d'examen dans l'en-tête.", "Connue : ré-identification prévue (FONC-005 en attente)"],
    ["Rapport : AB et BF_SH (cuisses) segmentés mais absents des tableaux de valeurs.", "Ouverte : mise en page à arbitrer avec le médecin"],
    ["Ligne de commande : une série mal saisie affiche « Internal error » au lieu d'un message clair.", "Ouverte"],
    ["Ligne de commande : un dossier DICOM inexistant affiche « Internal error » au lieu d'un message clair.", "Ouverte"],
], widths=[330, 130])

# --- 7. limites ------------------------------------------------------------
doc.add_heading("7. Limites et tests restant à faire", level=1)
bullets([
    "Équivalence T2 / mutools : non couverte (données de référence T2 à constituer).",
    "Cycle complet pause QC → reprise → résultat avec une vraie méthode : non couvert. "
    "Les étapes sont testées séparément ; l'interface graphique n'utilise pas ce chemin en phase 1.",
    "Interface graphique : non couverte par les tests automatisés (tests VAL avec les techniciens).",
    "Code mort identifié et déplacé dans bin/ (medical_report, dicom) : hors produit, non testé.",
    "dicom : échanges réels avec le PACS non automatisés ; l'option --anonymize des scripts C-MOVE n'est pas testée "
    "car elle n'est pas implémentée (incident ouvert).",
    "Rapport : le texte du PDF n'est pas relu directement (contrôle sur le HTML converti) ; la mise en page "
    "et les images (orientation des coupes) relèvent de RES-003.",
    "Équivalence Dixon : un seul examen et un seul segment (jambes).",
])

# --- 8. preuves ------------------------------------------------------------
doc.add_heading("8. Preuves d'exécution : le dossier resultats_tests", level=1)
doc.add_heading("À quel besoin il répond", level=2)
para("Un résultat affiché dans un terminal disparaît à sa fermeture. Le dossier resultats_tests conserve, "
     "pour chaque exécution, la sortie complète de pytest dans un fichier horodaté et relié à une version "
     "précise du code. Il permet à un tiers de vérifier chaque ligne du Journal de tests : ce qui a été "
     "exécuté, quand, sur quelle version, avec quel résultat.")
doc.add_heading("Où le trouver", level=2)
bullets([
    "Sur la station : /mnt/rmn_files/Users/IB/mureport/resultats_tests/ "
    "(depuis un PC : N:\\Users\\IB\\mureport\\resultats_tests\\).",
    "Copie de travail : PIPELINE_01\\resultats_tests\\.",
    "Le dossier n'est pas versionné dans git : il contient des traces d'exécution, pas du code.",
])
doc.add_heading("Contenu", level=2)
bullets([
    "Nommage : <ID du test>_<AAAAMMJJ>_<commit>, par exemple AUTO-001_20261022_37c3b0b.",
    ".txt : sortie lisible, une ligne par cas de test (PASSED, FAILED, XFAIL) et le bilan final.",
    ".xml : même résultat au format JUnit, standard, lisible par des outils.",
])
doc.add_heading("Reproduire une exécution", level=2)
para("Sur la station, dans /mnt/rmn_files/Users/IB/mureport, environnement « pipeline » actif et Docker lancé :")
for line in ("git checkout 6e33e57 && git submodule update",
             "pip install -e dicom",
             "F=resultats_tests/AUTO-001_$(date +%Y%m%d)_$(git rev-parse --short HEAD)",
             "python -m pytest tests/ medical_report/tests dicom/tests -v -rA --junitxml=$F.xml | tee $F.txt"):
    p = para(line, size=9.5, space_after=1)
    p.runs[0].font.name = "Consolas"

out = "MuReport-rapport-tests-AUTO-001.docx"
doc.save(out)
print("saved", out)


