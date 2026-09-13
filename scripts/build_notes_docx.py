"""Build a clean, structured Word document from the 2026-09-09 meeting notes."""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x26, 0x46, 0x53)      # deep teal
ACCENT_HEX = "264653"
LIGHT_HEX = "E8EEF0"                      # pale teal fill
CALLOUT_HEX = "F3EDE2"                    # warm sand fill
GREY = RGBColor(0x6B, 0x6B, 0x6B)
DARK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Pt(64)
sec.top_margin = sec.bottom_margin = Pt(56)

# --- base styles -----------------------------------------------------------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.12

for lvl, size in ((1, 15), (2, 11.5)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = ACCENT
    h.paragraph_format.space_before = Pt(18 if lvl == 1 else 10)
    h.paragraph_format.space_after = Pt(5)
    h.paragraph_format.keep_with_next = True

for s in ("List Bullet", "List Bullet 2", "List Number"):
    st = doc.styles[s]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = DARK
    st.paragraph_format.space_after = Pt(3)


def _shade(el, hex_fill):
    tcpr = el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def bottom_border(paragraph, hex_color=ACCENT_HEX, sz="8"):
    p = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), hex_color)
    borders.append(bottom)
    p.append(borders)


def para(text, size=None, color=None, italic=False, bold=False,
         align=None, space_before=None, space_after=None):
    par = doc.add_paragraph()
    run = par.add_run(text)
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    run.font.italic = italic
    run.font.bold = bold
    if align is not None:
        par.alignment = align
    if space_before is not None:
        par.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        par.paragraph_format.space_after = Pt(space_after)
    return par


def bullets(items, style="List Bullet"):
    for it in items:
        if isinstance(it, tuple):
            para_bullet(it[0], style)
            bullets(it[1], style="List Bullet 2")
        else:
            para_bullet(it, style)


def para_bullet(text, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def callout(label, text, fill=CALLOUT_HEX):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    _shade(cell._tc, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + " ")
    r.font.bold = True
    r.font.color.rgb = ACCENT
    r.font.size = Pt(10.5)
    r2 = p.add_run(text)
    r2.font.size = Pt(10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def data_table(headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    for i, h in enumerate(headers):
        _shade(hdr[i]._tc, ACCENT_HEX)
        p = hdr[i].paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        cells = tbl.add_row().cells
        for ci, val in enumerate(row):
            if ri % 2 == 1:
                _shade(cells[ci]._tc, LIGHT_HEX)
            p = cells[ci].paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(10)
            if ci == 0:
                r.font.bold = True
                r.font.color.rgb = ACCENT
    if widths:
        for row in tbl.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Pt(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def checklist(items):
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    for it in items:
        cells = tbl.add_row().cells
        cells[0].width = Pt(20)
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rb = p0.add_run("\u2610")
        rb.font.size = Pt(12)
        rb.font.color.rgb = ACCENT
        r = cells[1].paragraphs[0].add_run(it)
        r.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


# --- title ---------------------------------------------------------------
para("Notes de réunion", size=24, bold=True, color=ACCENT, space_after=2)
t = para("PYB · IB  —  9 septembre 2026", size=11, color=GREY, space_after=12)
bottom_border(t)
para("Projet PIPELINE_01 — points abordés et décisions.",
     italic=True, color=GREY, space_after=14)

# --- synthèse des décisions --------------------------------------------
doc.add_heading("Synthèse des décisions", level=1)
data_table(
    ["Sujet", "Décision"],
    [
        ["Multicentrique",
         "Mis en stand-by. Chantier à part entière : module dédié, pas un simple flag."],
        ["QC / checkpoint",
         "Les découpler au niveau des commandes. « resume » ne doit que reprendre là où on "
         "s'est arrêté."],
        ["QC en direct",
         "Le QC global est lancé par une commande dédiée ; le QC en direct passe par un "
         "formulaire (outil interactif)."],
        ["Entrée de « run »",
         "Option 2 : passer la stack DICOM, filtrage dans la méthode, avec un checkpoint "
         "dédié."],
        ["Étape 1 — report",
         "Retirer la génération de report à chaque run de « process » ; retirer le « pending » "
         "du QC pour la v1."],
    ],
    widths=[110, 360],
)

# --- 1. CLI ------------------------------------------------------------------
doc.add_heading("1.  Interface en ligne de commande", level=1)
bullets([
    "Déplacer les imports du haut des fichiers pour qu'ils ne soient chargés qu'au moment "
    "où c'est nécessaire : un simple « pipeline » ne doit pas être aussi lent.",
    "Ajouter une barre de progression.",
    "La ligne de commande doit refléter fidèlement la logique interne du programme.",
    "Une commande pour consulter : les méthodes disponibles, leurs checkpoints, les "
    "paramètres de chaque méthode ; et lister les jobs avec leur état.",
    "Option « segment » : la passer en mode choix (choice) pour éviter les fautes de frappe.",
])

# --- 2. Multicentrique -----------------------------------------------------
doc.add_heading("2.  Traitement des données multicentriques", level=1)
callout("Stand-by.",
        "Chantier à part entière : module dédié et vraie réflexion, bien plus qu'un simple flag.")
bullets([
    "Un reader spécifique.",
    "Du code dédié pour chaque site et chaque protocole.",
])

# --- 3. QC ---------------------------------------------------------------
doc.add_heading("3.  Contrôle qualité (QC)", level=1)

doc.add_heading("Affichage", level=2)
bullets([
    "Revoir l'affichage de l'overview en 5 × 4.",
    "Pour le GIF : changer l'orientation, et y associer le fichier de labels afin "
    "d'appliquer les bonnes couleurs.",
])

doc.add_heading("Concept", level=2)
callout("Priorité élevée.",
        "Découpler le QC du checkpoint au niveau de la commande. « resume » fait trop de "
        "choses ; il doit uniquement permettre de reprendre là où on s'est arrêté — et n'est "
        "valable que lorsqu'il n'y a aucune modification à faire.",
        fill="FBE9E7")
bullets([
    "Ajouter une commande « QC ».",
    "Via les jobs, « resume » dispose aujourd'hui des mêmes informations que « process ». "
    "Il faut que « process » ait les mêmes informations et puisse faire les mêmes actions. "
    "Exemple : ne pas reprendre mais tout refaire de A à Z avec les changements nécessaires, "
    "depuis « process ».",
    "« action » doit devenir une sorte de « paramètre ».",
    "Ajouter un flag « reset » : revenir à zéro, ou reprendre à un checkpoint précis en le "
    "nommant.",
])
callout("Décision.",
        "Le QC global est lancé par une commande. Pour un QC en direct, on remplit un "
        "formulaire pour la décision : l'outil devient interactif.")
bullets([
    "À réfléchir : comment concilier cela avec l'architecture client / serveur. Le QC "
    "correspondra à terme à des éléments plus complexes de la base de données.",
])

# --- 4. Input de run -----------------------------------------------------
doc.add_heading("4.  Entrée de « run »", level=1)
para("Actuellement — point qui n'avait pas vraiment été pensé — on passe directement les "
     "DICOM (stack brute) à la fonction « run » de la méthode.")
data_table(
    ["Option", "Principe", "Conséquence / besoin"],
    [
        ["1 — Actuel",
         "La méthode appelle DicomStack puis filtre les séries.",
         "Besoin : numéro de série, date, study."],
        ["2 — Stack",
         "L'entrée est la stack ; la méthode ne fait que le filtrage.",
         "On peut reprendre depuis la stack DICOM avec les bons numéros de série."],
        ["3 — Stack filtrée",
         "L'entrée est la stack déjà filtrée.",
         "La méthode n'a besoin de rien."],
    ],
    widths=[85, 200, 185],
)
callout("Décision.",
        "Viser l'option 2 (DICOM + filtrage dans la méthode), en ajoutant un checkpoint.")
bullets([
    "Les paramètres de filtrage doivent être modifiables par l'utilisateur : la logique "
    "veut que je puisse changer les séries.",
    "Éviter de retélécharger les DICOM à distance et de ressaisir tous les autres "
    "paramètres → « resume reset ». Refaire l'ensemble reste possible via « process ».",
    "Que faire si un jour on a des données qui ne sont pas des DICOM ? L'idée : lancer le "
    "téléchargement une seule fois, puis les 4 commandes ; éviter de re-télécharger grâce "
    "aux checkpoints.",
])

# --- 5. T2 -------------------------------------------------------------------
doc.add_heading("5.  T2", level=1)
bullets([
    "Faire un overview sur les T2 maps : on peut passer l'overview sur tout un volume. "
    "Avoir les deux volumes, T2 et ffmap, dans le même overview pour croiser les informations.",
    "mutools génère aussi des courbes de fit : visualisation du fit sur différentes régions "
    "de l'image.",
    "Le clustering = 3 éléments : background, fat, water.",
    "On peut générer une image pour le QC.",
    "« mutools t2map create… » génère cette image. Le faire et évaluer sa pertinence, en "
    "prévision du multicentrique.",
    "Sur l'overview : possibilité d'ajouter une colorbar et de colorer pour le QC "
    "multicentrique.",
])

# --- 6. Étape 1 --------------------------------------------------------------
doc.add_heading("6.  Étape 1 — à faire", level=1)
checklist([
    "Extraire le QC.",
    "« process » avec global_swap.",
    "DICOM sur disque pour l'instant.",
    "Changements sur « action ».",
    "« resume ».",
    "V1 : enlever le « pending » du QC pour simplifier.",
    "Enlever la génération de report à chaque run de « process ».",
    "Réfléchir à la manière de tout lancer d'un coup.",
    "Faire une GUI.",
    "Savoir à l'avance quels sont les bons numéros de série.",
    "Permettre à l'utilisateur de voir le contenu de l'exam, trier et trouver les bons "
    "numéros de série (viewer type RadiAnt ? affichage du contenu de l'exam).",
])

out = r"c:\Users\i.benziane\Documents\PIPELINE_01\Notes_reunion_2026-09-09.docx"
doc.save(out)
print("saved:", out)
